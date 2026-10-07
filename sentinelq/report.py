"""Stage 6 - scorecard + evidence trail for every score + coverage sheet + run log."""
from __future__ import annotations
import csv
import json
from dataclasses import asdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import parse_date

HDR = PatternFill("solid", fgColor="1F3864")
FILLS = {"Clean": "C6EFCE", "Watch": "FFEB9C", "Flag": "FFC7CE"}


def _sheet(wb, title, headers, rows, link_col=None, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HDR
    for r in rows:
        ws.append(r)
    if link_col is not None:
        for row in ws.iter_rows(min_row=2, min_col=link_col + 1, max_col=link_col + 1):
            for c in row:
                if c.value and str(c.value).startswith("http"):
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, len(h) + 2))
    ws.freeze_panes = "A2"
    return ws


def save_or_sidestep(write, path: Path) -> Path:
    """Write a report file; if it is open in Excel / a PDF viewer (Windows locks it), write a timestamped copy instead of
    failing the whole run at its last step. Returns the path actually written."""
    try:
        write(path)
        return path
    except PermissionError:
        from datetime import datetime
        alt = path.with_name(f"{path.stem}_{datetime.now():%H%M%S}{path.suffix}")
        write(alt)
        print(f"  WARNING: {path.name} is open in another program (locked) - wrote {alt.name} instead. Close the file before the next run.")
        return alt


def _gist_line(pos, neg) -> str:
    """Fallback one-line read: the event kinds that moved the score most each way."""
    from .score import _PLAIN
    kinds = lambda xs: ", ".join(dict.fromkeys(_PLAIN.get(li.label.event_type, li.label.event_type.replace("_", " ")) for li, _ in xs[:3]))
    parts = ([f"Up: {kinds(pos)}"] if pos else []) + ([f"Down: {kinds(neg)}"] if neg else [])
    return " | ".join(parts) or "No signal either way"


def write_reports(out: Path, res: dict, rubric) -> None:
    out.mkdir(parents=True, exist_ok=True)
    scores, kept, run = res["scores"], res["kept"], res["run"]
    wb = Workbook()
    wb.remove(wb.active)

    # Scorecard - whole numbers only; the 'why' is in words
    from datetime import date
    from .config import SAMPLE_WINDOWS
    from .score import event_treatment
    as_of = date.fromisoformat(str(run.get("as_of"))[:10]) if run.get("as_of") else date.today()
    treat = {}
    for v in kept.values():
        treat.update(event_treatment(v, as_of, rubric))

    def win(li):
        if li.item.window:
            return int(li.item.window)
        d = parse_date(li.item.published)
        age = (as_of - d).days if d else 0
        return next((i for i, (lo, hi, _) in enumerate(SAMPLE_WINDOWS, 1) if lo <= age <= hi), len(SAMPLE_WINDOWS))
    wname = {i: f"W{i} ({lo}-{hi}d)" for i, (lo, hi, _) in enumerate(SAMPLE_WINDOWS, 1)}
    arts = {sym: [li for li in v if li.item.kind == "news"] for sym, v in kept.items()}

    from .score import driver_line, drivers
    b_ = rubric["sentiment"].get("bands") or {"plus2": 55, "plus1": 20}
    band_txt = {2: f"+{b_['plus2']} or more", 1: f"+{b_['plus1']} to +{b_['plus2'] - 1}", 0: f"-{b_['plus1'] - 1} to +{b_['plus1'] - 1}",
                -1: f"-{b_['plus2'] - 1} to -{b_['plus1']}", -2: f"-{b_['plus2']} or less"}
    narr = res.get("narrative") or {}
    heads = ["Symbol", "Company", "Sentiment (-2..+2)", "Governance Score", "Governance Label", "In one line", "Sentiment - why",
             "Main positives", "Main negatives", "Governance - why", "Net evidence (-100..+100)", "Cap", "Sector", "Weight",
             "Sector Sentiment (-2..+2)", "Corporate Action Score", "Articles scored", "Full text read",
             "Articles by window (W1 / W2 / W3 / W4)", "Low Confidence"]
    rows = []
    for s_ in scores:
        a_ = arts.get(s_.symbol, [])
        wc = [sum(1 for li in a_ if win(li) == k) for k in (1, 2, 3, 4)]
        pos, neg, f_ = drivers(kept.get(s_.symbol, []), as_of, rubric)
        if s_.company_sentiment is None:
            swhy = "Not scored - " + (s_.sentiment_note.replace("INSUFFICIENT DATA: ", "") or "insufficient evidence")
        else:
            n_zero = sum(1 for li in kept.get(s_.symbol, []) if treat.get(id(li), {}).get("counted") == "yes" and li.label.sentiment == 0)
            swhy = (f"Average {f_['points']:+d} falls in the {s_.company_sentiment:+d} band ({band_txt[s_.company_sentiment]}). "
                    f"{f_['n_pos']} positive, {n_zero} neutral, {f_['n_neg']} negative articles; last 3 months carry {f_['recent']:.0%} of the weight.")
        pens = s_.governance_penalties
        if pens:
            gwhy = "\n".join(f"{p_['penalty']:+g} {p_['event_type'].replace('_', ' ')} ({p_['date']}): {p_['headline'][:70]}" for p_ in pens)
        else:
            gwhy = "No penalised governance event in the 12 months."
        if s_.governance_ignored:
            gwhy += f"\n{len(s_.governance_ignored)} flagged item(s) checked and not penalised - see Evidence - Governance."
        rows.append([s_.symbol, s_.name, s_.company_sentiment, s_.governance_score, s_.governance_label,
                     (narr.get(s_.symbol) or {}).get("one_line_read") or _gist_line(pos, neg), swhy,
                     "\n".join(driver_line(li, p) for li, p in pos[:3]) or "none", "\n".join(driver_line(li, p) for li, p in neg[:3]) or "none",
                     gwhy, s_.sentiment_points, s_.cap, s_.sector, s_.weight, s_.sector_sentiment, round(s_.corporate_action_score), len(a_),
                     sum(1 for li in a_ if (li.item.parse or "").startswith("full-text")), " / ".join(str(x) for x in wc),
                     "YES" if s_.low_confidence else ""])
    ws = _sheet(wb, "Scorecard", heads, rows, widths={"Company": 22, "Sentiment (-2..+2)": 11, "Governance Score": 11, "Governance Label": 11,
                                                      "In one line": 26, "Sentiment - why": 42, "Main positives": 66, "Main negatives": 66,
                                                      "Governance - why": 48, "Articles by window (W1 / W2 / W3 / W4)": 18})
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 32
    sfill = {2: "C6EFCE", 1: "E2F0D9", 0: "FFFFFF", -1: "FCE4D6", -2: "FFC7CE"}
    for c in ws[1]:
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 62
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[2].fill = PatternFill("solid", fgColor=sfill.get(row[2].value, "FFFFFF"))
        row[4].fill = PatternFill("solid", fgColor=FILLS.get(row[4].value, "FFFFFF"))
        for c in (0, 2, 3, 4, 5):
            row[c].font = Font(bold=True)
        for c in (2, 3, 4):
            row[c].alignment = Alignment(horizontal="center", vertical="top")

    rm = wb.create_sheet("Read me", 0)
    guide = [("How to read the Scorecard", ""),
             ("Layout", "Left to right: the scores, a one-line read, WHY (sentiment, the articles that moved it most, governance), then the detail."),
             ("In one line", "A short plain-English read of the stock's news year (from the IC commentary)."),
             ("Sentiment (-2..+2)", "The stock's news score. Every article about the company in the 12-month sample is read and scored -2..+2."),
             ("Net evidence (-100..+100)", "The weighted average of those article scores x 50. Weights: W1 (0-30 days) 1.0, W2 (31-90) 0.75, "
                                           "W3 (91-180) 0.5, W4 (181-365) 0.25. Neutral (0) articles count and pull toward 0. It decides the band below."),
             ("Bands", f"+2: {band_txt[2]}   |   +1: {band_txt[1]}   |   0: {band_txt[0]}   |   -1: {band_txt[-1]}   |   -2: {band_txt[-2]}"),
             ("Sentiment - why", "Which band the net falls in, how many events pushed up vs down, and how much of the weight is recent."),
             ("Main positives / negatives", "The three articles that moved the score most each way: article score | date | headline (points it added)."),
             ("Governance Score", "Starts at 100; fixed penalties for verified adverse events (regulator orders, auditor exits, CXO exits...). "
                                  "Penalties older than 6 months fade to half by 12 months."),
             ("Governance Label", "Clean 95+  |  Watch 80-94  |  Flag below 80. Thresholds unchanged."),
             ("Governance - why", "Each penalty applied, with date and headline; items checked but not penalised are listed in 'Evidence - Governance'."),
             ("Articles by window", "Articles scored in W1 0-30 days / W2 31-90 / W3 91-180 / W4 181-365 (sample targets 40/25/20/15)."),
             ("", ""),
             ("Other sheets", ""),
             ("Articles by window", "Every sampled article: what it says, its score and why, and what it did to the stock's score (filter by Symbol / Window)."),
             ("Window summary", "Per stock and window: articles scored, sampled-but-not-scored, the +2..-2 counts and the points each window added."),
             ("Evidence - Governance", "Every governance item: penalty applied, or why it was not penalised.")]
    for r_ in guide:
        rm.append(list(r_))
    rm.column_dimensions["A"].width = 28
    rm.column_dimensions["B"].width = 130
    for row in rm.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    rm["A1"].font = Font(bold=True, size=13)
    rm["A12"].font = Font(bold=True, size=13)

    # Every article, by window: what it says, how it was scored and why, and what it did to the stock's score
    pen = {}
    for s_ in scores:
        for p_ in s_.governance_penalties:
            pen[(s_.symbol, p_.get("url"))] = f"penalty {p_['penalty']:+g} ({p_['event_type'].replace('_', ' ')})"
        for p_ in s_.governance_ignored:
            pen.setdefault((s_.symbol, p_.get("url")), "flagged, not penalised: " + p_["why"])
    names = {s_.symbol: s_.name for s_ in scores}
    art_head = ["Symbol", "Company", "Window", "Date", "Age (days)", "Headline", "Source", "Read from", "What the article says",
                "Event type", "Article score (-2..+2)", "Why this score", "Counted in sentiment", "Window weight (%)",
                "Points added to stock average", "How it was treated", "Governance", "Also reported by", "Source URL"]
    art_rows = []
    for sym in sorted(arts):
        for li in sorted(arts[sym], key=lambda x: (win(x), -(parse_date(x.item.published) or as_of).toordinal())):
            it, l = li.item, li.label
            t = treat.get(id(li), {"counted": "no", "share": 0, "wweight": "", "points": "", "why": ""})
            d = parse_date(it.published)
            art_rows.append([sym, names.get(sym, ""), wname[win(li)], it.published, (as_of - d).days if d else "", it.title, it.source,
                             "full text" if (it.parse or "").startswith("full-text") else "headline only", l.gist or "",
                             l.event_type, l.sentiment, l.rationale, t["counted"], t.get("wweight", ""), t["points"], t["why"],
                             pen.get((sym, it.url), "flagged" if l.governance_flag else ""),
                             len(it.members or []) or "", it.url])
    # Articles that were in the window sample but did not reach the score: shown too, with the plain reason (and the model's read when it read one)
    why_drop = {"model_boilerplate": "templated piece with no company event (judged boilerplate on reading)",
                "not_about_company": "not about this company (e.g. a different company with a similar name, or a list)",
                "label_failed": "the model could not label it",
                "results_quota_one_print_per_quarter": "another article on the same quarter's results was kept (one results print per quarter)",
                "results_quota_other_per_quarter": "over the per-quarter limit for other results-type stories",
                "per_day_cap_no_same_day_cluster": "over the per-day cap for that date"}
    rawlab = {}
    for r_ in res.get("raw_log") or []:
        try:
            j = json.loads(r_.get("raw") or "")
            if isinstance(j, dict) and "event_type" in j:
                rawlab[(r_.get("symbol"), r_.get("url"))] = j
        except (ValueError, TypeError):
            pass
    for d_ in res.get("dropped") or []:
        code = (d_.reason or "").split(":")[0]
        if code not in why_drop or d_.symbol not in names:
            continue
        j = rawlab.get((d_.symbol, d_.url), {})
        dd = parse_date(d_.date)
        age = (as_of - dd).days if dd else None
        k = next((i for i, (lo, hi, _) in enumerate(SAMPLE_WINDOWS, 1) if age is not None and lo <= age <= hi), len(SAMPLE_WINDOWS))
        art_rows.append([d_.symbol, names.get(d_.symbol, ""), wname[k], d_.date, age if age is not None else "", d_.headline, "",
                         "read by model" if j else "not read", j.get("gist", ""), j.get("event_type", ""), j.get("sentiment", ""),
                         j.get("rationale", ""), "no", "", "", "not scored: " + why_drop[code], "", "", d_.url])
    art_rows.sort(key=lambda r_: (r_[0], r_[2], -(parse_date(r_[3]) or as_of).toordinal()))

    W = {"Headline": 60, "What the article says": 80, "Why this score": 80, "How it was treated": 55, "Source URL": 40, "Event type": 20,
         "Governance": 40, "Company": 24, "Window": 12}
    ws = _sheet(wb, "Articles by window", art_head, art_rows, link_col=18, widths=W)
    ws.auto_filter.ref = ws.dimensions
    sc_fill = {2: "C6EFCE", 1: "E2F0D9", 0: "FFFFFF", -1: "FCE4D6", -2: "FFC7CE"}
    for row in ws.iter_rows(min_row=2):
        row[10].fill = PatternFill("solid", fgColor=sc_fill.get(row[10].value, "FFFFFF"))
        for c in (5, 8, 11, 15):
            row[c].alignment = Alignment(wrap_text=True, vertical="top")

    # Per stock x window: what each period's news said, for human validation of the time periods
    summ = []
    for sym in sorted(arts):
        for k in (1, 2, 3, 4):
            a_ = [li for li in arts[sym] if win(li) == k]
            cnt = {v: sum(1 for li in a_ if li.label.sentiment == v) for v in (2, 1, 0, -1, -2)}
            n_out = sum(1 for r_ in art_rows if r_[0] == sym and r_[2] == wname[k] and str(r_[15]).startswith("not scored"))
            summ.append([sym, names.get(sym, ""), wname[k], len(a_), n_out, cnt[2], cnt[1], cnt[0], cnt[-1], cnt[-2],
                         sum(treat.get(id(li), {}).get("share", 0) for li in a_), sum(treat.get(id(li), {}).get("points") or 0 for li in a_),
                         sum(1 for li in a_ if (li.item.parse or "").startswith("full-text"))])
    ws = _sheet(wb, "Window summary", ["Symbol", "Company", "Window", "Articles scored", "Sampled, not scored", "+2", "+1", "0", "-1", "-2", "Window's share of stock weight (%)",
                                       "Points added to stock average", "Full text read"], summ, widths={"Company": 24, "Window": 12})
    ws.auto_filter.ref = ws.dimensions

    def vrec(p):
        v = p.get("verification") or {}
        return (f"subject_is_company={v.get('subject_is_company')}; event_at_company={v.get('event_at_company')}; direction={v.get('direction')}; "
                f"severity={v.get('severity')}; stage={v.get('stage')}; amount_cr={v.get('amount_inr_cr')}") if v else ""
    gov = [[s.symbol, p["date"], p["headline"], p["event_type"], p["penalty"],
            "scored - " + p.get("basis", "") + (f" [{p['verified']}]" if p.get("verified") else ""), p["url"] or p.get("source_ref", ""),
            vrec(p), p.get("n_sources"), p.get("max_tier"), p.get("confidence")]
           for s in scores for p in s.governance_penalties]
    gov += [[s.symbol, p["date"], p["headline"], p["event_type"], 0, "NOT scored: " + p["why"], p["url"] or p.get("source_ref", ""), "", "", "", ""]
            for s in scores for p in s.governance_ignored]
    _sheet(wb, "Evidence - Governance", ["Symbol", "Date", "Headline", "Event Type", "Penalty", "Treatment", "Source URL", "Verification record",
                                         "n_sources", "max_tier", "Confidence"], gov, link_col=6, widths={**W, "Treatment": 55, "Verification record": 70})

    ca = [[s.symbol, d["date"], d["headline"], d["event_type"], d["materiality"], d["contribution"], d["url"]]
          for s in scores for d in s.corporate_action_detail]
    _sheet(wb, "Evidence - Corp Actions",
           ["Symbol", "Date", "Headline", "Event Type", "Materiality", "Contribution", "Source URL"],
           ca, link_col=6, widths=W)

    cmap = {s.symbol: s for s in scores}
    cov = [[c["symbol"], c["retrieved"], c["news_kept"], c["actions_kept"], c["kept"], c["dropped"], c["label_failed"],
            c.get("relevant_articles", ""), c.get("span_days", ""), c.get("sentiment_status", ""), "YES" if c["low_confidence"] else "",
            cmap[c["symbol"]].coverage_map if c["symbol"] in cmap else "", "YES" if c["symbol"] in cmap and cmap[c["symbol"]].window_dominated else ""]
           for c in res["coverage"]]
    t = run["totals"]
    cov.append(["PORTFOLIO", t["retrieved"], "", "", t["kept"], t["dropped"], t["label_failed"], "", "", "", "", "", ""])
    _sheet(wb, "Coverage", ["Symbol", "Retrieved", "News Kept", "Actions Kept", "Kept", "Dropped",
                            "Of which: model failed to label", "Relevant articles", "Evidence span (days)", "Sentiment status", "Low Confidence",
                            "Coverage map", "One window dominates (>70%)"], cov, widths={"Sentiment status": 60, "Coverage map": 70})
    ws = wb["Coverage"]
    ws.append([])
    ws.append(["Dropped items (stage, reason)"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for d in res["dropped"]:
        ws.append([d.symbol, d.date, d.headline, d.stage, d.reason, d.url])
    for e in res["ingest_errors"]:
        ws.append([e["symbol"], "INGEST ERROR", e["provider"], e["error"]])

    if res["price_context"]:
        pc = res["price_context"]
        keys = ["symbol", "ret_1m", "ret_3m", "ret_6m", "excess_1m", "realized_vol_ann", "quadrant"]
        _sheet(wb, "Price Context (info only)",
               ["Symbol", "1M Ret", "3M Ret", "6M Ret", "1M Excess vs Nifty", "Realized Vol (ann.)",
                "Sentiment x Excess Quadrant"],
               [[p[k] if not isinstance(p[k], float) else round(p[k], 4) for k in keys] for p in pc],
               widths={"Sentiment x Excess Quadrant": 50})
    if res["relative"]:
        keys = list(res["relative"][0])
        _sheet(wb, "Relative (z-scores)", keys, [[r[k] for k in keys] for r in res["relative"]])

    delta = res.get("delta")
    if delta:                                   # B6: computed from scores/ across the last two run_ids, never re-derived from news
        _sheet(wb, "Delta", ["Symbol", "Prev run", "Prev sentiment", "Prev governance", "Prev label", "This sentiment", "This governance", "This label", "Change"],
               [[d["symbol"], d["prev_run"], d["prev_sentiment"], d["prev_governance"], d["prev_label"], d["sentiment"], d["governance"], d["label"], d["change"]]
                for d in delta], widths={"Prev run": 28, "Change": 50})
    _sheet(wb, "Run Info", ["Key", "Value"], [[k, json.dumps(v) if isinstance(v, (dict, list)) else v]
                                              for k, v in run.items()], widths={"Value": 120})
    save_or_sidestep(lambda f: wb.save(f), out / "sentinelq_report.xlsx")

    # Machine-readable outputs
    with open(out / "scorecard.csv", "w", newline="", encoding="utf-8-sig") as f:   # headlines (Rs sign etc.) appear in the rationale
        w = csv.writer(f)
        w.writerow(heads)
        w.writerows(rows)
    (out / "scores.json").write_text(json.dumps([s.to_dict() for s in scores], indent=2))
    (out / "evidence.json").write_text(json.dumps(
        [li.to_row() for v in kept.values() for li in v], indent=2))
    # Audit trail
    (out / "run_meta.json").write_text(json.dumps(
        dict(run, coverage=res["coverage"], ingest_errors=res["ingest_errors"],
             rubric=rubric.data), indent=2))
    with open(out / "dropped.jsonl", "w", encoding="utf-8") as f:
        for d in res["dropped"]:
            f.write(json.dumps(asdict(d)) + "\n")
    with open(out / "raw_model_responses.jsonl", "w", encoding="utf-8") as f:
        for r in res["raw_log"]:
            f.write(json.dumps(r) + "\n")
