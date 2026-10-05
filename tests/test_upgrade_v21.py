"""Upgrade Workflow v2.1 - acceptance tests A5 (Change A) and B7 (Change B) plus unit tests of the new modules."""
import json
from datetime import date, timedelta
from pathlib import Path

import pytest

import sentinelq.config as cfg
from sentinelq import cluster as CL
from sentinelq import replay as RP
from sentinelq.classify import KeywordClassifier
from sentinelq.corpus import Corpus
from sentinelq.golden import run_golden
from sentinelq.models import Holding, RawItem
from sentinelq.pipeline import Pipeline
from sentinelq.rubric import load_rubric
from sentinelq.sampler import stratified_sample
from sentinelq.tiers import domain_of, tier_of

ROOT = Path(__file__).resolve().parent.parent
AS_OF = date(2026, 7, 3)
R = load_rubric()
DOMS = ["reuters.com", "economictimes.indiatimes.com", "livemint.com", "business-standard.com", "moneycontrol.com", "thehindu.com",
        "cnbctv18.com", "ndtvprofit.com", "financialexpress.com", "thehindubusinessline.com", "zeebiz.com", "goodreturns.in",
        "ptinews.com", "bloomberg.com", "lokmattimes.com", "thehansindia.com"]


def art(sym, title, d, i, dom=None, snippet=""):
    dom = dom or DOMS[i % len(DOMS)]
    return RawItem(sym, "news", title, snippet, f"https://{dom}/{sym}/{d}/{i}", d.isoformat(), dom)


class Fixed:
    def __init__(self, items):
        self.items = items
        self.audit = []

    def fetch(self, h, s, e):
        return [i for i in self.items if i.symbol == h.symbol and s <= date.fromisoformat(i.published) <= e]


class CountingClassifier(KeywordClassifier):
    calls = 0

    def classify(self, item, company):
        CountingClassifier.calls += 1
        return super().classify(item, company)


def bajaj_auto_year():
    """A year of Bajaj Auto news: steady operating updates in every window, a results print per quarter, and 40 outlets on the
    24-26 Jun ransomware story (same-day flood)."""
    sym, items = "BAJAJ-AUTO", []
    for k in range(0, 365, 9):
        d = AS_OF - timedelta(days=k)
        items.append(art(sym, f"Bajaj Auto business update {d}: export volumes yoy growth", d, k))
        items.append(art(sym, f"Bajaj Auto reports export volumes yoy growth in business update {d}", d, k + 2000))   # same event, other wording
    for q, days in enumerate((20, 110, 200, 290)):
        d = AS_OF - timedelta(days=days)
        items.append(art(sym, f"Bajaj Auto Q{4 - q} results: profit rises on strong exports", d, 500 + q))
    for i in range(40):
        d = AS_OF - timedelta(days=7 + (i % 3))
        items.append(art(sym, f"Bajaj Auto says operations uninterrupted as ransomware investigation continues - outlet {i}", d, 900 + i, f"outlet{i}.com"))
    return items


def run(tmp_path, items, holdings, stratified=True, corpus=None, clf=None, name="o", **kw):
    p = Pipeline(R, Fixed(items), None, None, clf or KeywordClassifier(), tmp_path / name, None, AS_OF, governance_pass=False,
                 fundamentals_pass=False, stratified=stratified, corpus_dir=corpus, disk_cache=False, **kw)
    return p.run(holdings)


H = [Holding("BAJAJ-AUTO", "Bajaj Auto", "Autos", "Large")]


# ============================================================== A1 / A2
def test_a1_stratified_quotas_and_w1_never_fills_the_budget():
    items = [art("X", f"X Corp update {k}", AS_OF - timedelta(days=k % 365), k) for k in range(400)]
    sel, extras, rep = stratified_sample(items, AS_OF, ["x corp"])
    assert rep["windows"] == {1: 40, 2: 25, 3: 20, 4: 15} and len(sel) == 100
    assert all(i.window for i in sel)


def test_a1_carry_forward_rolls_to_older_then_back_to_newest():
    items = [art("X", f"X Corp update {k}", AS_OF - timedelta(days=k), k) for k in range(0, 120)]       # nothing older than 120 days
    sel, extras, rep = stratified_sample(items, AS_OF, ["x corp"])
    # W1 holds 31 (short by 9 -> rolls to W2: 25+9), W3 holds 29 (takes its 20), W4 is empty (its 15 roll back to the newest window with
    # articles left: W2). One article a day, so the per-day cap never binds. Total budget stays 100.
    assert rep["windows"] == {1: 31, 2: 49, 3: 20, 4: 0} and len(sel) == 100
    assert rep["carry"][2].get("rolled_back") == 15


def test_a1_per_day_cap_six_rest_become_extra_sources():
    d = AS_OF - timedelta(days=3)
    items = [art("X", f"X Corp flood story {i}", d, i, f"o{i}.com") for i in range(20)] + [art("X", f"X Corp other {k}", AS_OF - timedelta(days=40 + k), k) for k in range(10)]
    sel, extras, rep = stratified_sample(items, AS_OF, ["x corp"])
    assert sum(1 for i in sel if i.published == d.isoformat()) == 6 and len(extras) == 14 and all(x.sample_extra for x in extras)


def test_a2_tiers_are_deterministic_by_domain_and_google_links_use_the_source():
    assert tier_of("reuters.com") == "T1" and tier_of("m.economictimes.indiatimes.com") == "T2" and tier_of("some-unknown-blog.in") == "T3"
    assert tier_of("marketscreener.com") == "T4"
    assert domain_of("https://news.google.com/rss/articles/CBMi", "livemint.com") == "livemint.com"


def test_a3_confidence_formula_matches_the_document():
    assert CL.confidence(1, "T1") == 0.504 and CL.confidence(11, "T1") == 0.773 and CL.confidence(1, "T4") == 0.252
    assert CL.confidence(1, "T3") < cfg.CONF_UNPENALISED_BELOW


# ============================================================== A5 acceptance
def test_a5_1_no_single_window_dominates(tmp_path):
    res = run(tmp_path, bajaj_auto_year(), H)
    s = res["scores"][0]
    assert s.n_events >= 6 and max(s.window_events.values()) / s.n_events <= 0.70, s.coverage_map
    assert not s.window_dominated and s.coverage_map.startswith("W1:")


def test_a5_2_ransomware_flood_resolves_to_one_governance_event_with_many_sources(tmp_path):
    res = run(tmp_path, bajaj_auto_year(), H)
    s = res["scores"][0]
    inv = [p for p in s.governance_penalties if p["event_type"] == "investigation"]
    assert len(inv) == 1 and inv[0]["n_sources"] >= 10, (inv, s.governance_ignored)
    ev = [li for li in res["kept"]["BAJAJ-AUTO"] if li.label.event_type == "investigation"]
    assert len(ev) == 1 and ev[0].item.n_members >= 10


def test_a5_3_notice_then_probe_resolve_to_one_governance_event(tmp_path):
    sym, items = "NESTLEIND", []
    for k in range(0, 365, 12):
        items.append(art(sym, f"Nestle India business update {k}: volume growth yoy", AS_OF - timedelta(days=k), k))
    items.append(art(sym, "Nestle India gets FSSAI notice over alleged Maggi contamination; probe launched", date(2026, 6, 14), 700, "reuters.com"))
    items.append(art(sym, "Nestle India faces FSSAI probe over Maggi quality concerns", date(2026, 6, 15), 701, "livemint.com"))
    res = run(tmp_path, items, [Holding("NESTLEIND", "Nestle India", "FMCG")])
    s = res["scores"][0]
    assert len([p for p in s.governance_penalties if p["event_type"] == "investigation"]) == 1
    assert s.governance_score == 90


def test_a5_4_results_events_anchor_sentiment_and_absence_is_reported(tmp_path, monkeypatch):
    monkeypatch.setattr(cfg, "MIN_EVENTS", 6)
    monkeypatch.setattr(cfg, "MIN_WINDOWS", 2)
    monkeypatch.setattr(cfg, "MIN_RESULTS_EVENTS", 1)
    res = run(tmp_path, bajaj_auto_year(), H, name="with")
    s = res["scores"][0]
    assert s.n_results_events >= 1 and s.company_sentiment is not None
    no_results = [i for i in bajaj_auto_year() if "results" not in i.title]
    for i in no_results:                                            # nothing the taxonomy reads as a results print
        i.title = i.title.replace("business update", "dealer note").replace("yoy growth", "steady")
    s2 = run(tmp_path, no_results, H, name="without")["scores"][0]
    assert s2.company_sentiment is None and "no results event" in s2.sentiment_note


def test_a5_5_sampler_and_clustering_are_deterministic():
    items = bajaj_auto_year()
    alias = CL.alias_tokens("Bajaj Auto", "", "BAJAJ-AUTO")

    def once():
        sel, extras, _ = stratified_sample([RawItem(**{**i.__dict__}) for i in items], AS_OF, ["bajaj auto"])
        common = CL.top_frequency([i.title for i in sel], alias)
        cl = CL.precluster(sel, alias, common)
        evs = [CL.build_event(c, "sentiment") for c in cl]
        return [(e.event_id, e.url, e.n_members) for e in evs]
    assert once() == once()


def test_a5_6_llm_calls_fall_by_half_versus_the_latest_n_sampler(tmp_path):
    items = bajaj_auto_year()
    CountingClassifier.calls = 0
    run(tmp_path, [RawItem(**{**i.__dict__}) for i in items], H, stratified=False, clf=CountingClassifier(), name="legacy", max_articles=100)
    legacy = CountingClassifier.calls
    CountingClassifier.calls = 0
    run(tmp_path, [RawItem(**{**i.__dict__}) for i in items], H, stratified=True, clf=CountingClassifier(), name="strat")
    strat = CountingClassifier.calls
    assert strat <= legacy * 0.5, (legacy, strat)


# ============================================================== B7 acceptance
def test_b7_1_run_then_replay_is_identical(tmp_path):
    res = run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus")
    rep = RP.replay(AS_OF, None, tmp_path / "corpus", tmp_path / "replay", verified=None, quiet=True)
    a = json.loads((tmp_path / "o" / "scores.json").read_text())
    b = json.loads((tmp_path / "replay" / "scores.json").read_text())
    assert a == b and rep["mode"] == "pit-replay"


def test_b7_2_replay_with_a_changed_weight_moves_only_rows_with_that_penalty(tmp_path):
    items = bajaj_auto_year() + [art("TITAN", f"Titan Company update {k}: jewellery sales yoy", AS_OF - timedelta(days=k), k) for k in range(0, 365, 15)]
    hs = H + [Holding("TITAN", "Titan Company", "Consumption")]
    run(tmp_path, items, hs, corpus=tmp_path / "corpus")
    rub = json.loads((ROOT / "rubric" / "rubric_v1.json").read_text())
    rub["governance"]["penalties"]["investigation"] = -30
    rub["version"] = "test-weight"
    (tmp_path / "rubric_x.json").write_text(json.dumps(rub))
    base = {s.symbol: s for s in RP.replay(AS_OF, None, tmp_path / "corpus", tmp_path / "r0", verified=None, quiet=True)["scores"]}
    chg = {s.symbol: s for s in RP.replay(AS_OF, str(tmp_path / "rubric_x.json"), tmp_path / "corpus", tmp_path / "r1", verified=None, quiet=True)["scores"]}
    assert chg["BAJAJ-AUTO"].governance_score == base["BAJAJ-AUTO"].governance_score - 20       # -10 -> -30
    assert chg["TITAN"].governance_score == base["TITAN"].governance_score and chg["TITAN"].company_sentiment_raw == base["TITAN"].company_sentiment_raw


def test_b7_3_backtest_refuses_articles_fetched_after_as_of_plus_7(tmp_path):
    run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus")
    c = Corpus(tmp_path / "corpus")
    m = c.manifests()[0]
    # a genuine point-in-time replay needs the run to have been FETCHED within as_of + 7 days; this run was fetched today (long after 3-Jul-2026)
    assert m["started_at"][:10] > (AS_OF + timedelta(days=7)).isoformat()
    rows = RP.backtest(AS_OF, AS_OF + timedelta(days=7), tmp_path / "corpus", tmp_path / "bt", quiet=True)
    assert all(r["mode"] == "no-data" for r in rows)
    # rewrite the manifest's fetch time into the grace window (the only way to simulate a live weekly run) -> visible, pit-replay
    m["started_at"] = AS_OF.isoformat() + "T18:00:00"
    (tmp_path / "corpus" / "manifest" / f"{m['run_id']}.json").write_text(json.dumps(m))
    rows = RP.backtest(AS_OF, AS_OF + timedelta(days=7), tmp_path / "corpus", tmp_path / "bt2", quiet=True)
    assert {r["mode"] for r in rows} == {"pit-replay"} and (tmp_path / "bt2" / "backtest.csv").exists()


def test_b7_4_golden_gate_blocks_publication(tmp_path):
    bad = tmp_path / "golden.jsonl"
    bad.write_text(json.dumps({"id": "X1", "symbol": "X", "headline": "X Corp appoints new CEO", "snippet": "", "published_date": "2026-06-01",
                               "source_domain": "reuters.com", "expected": "penalise", "expected_flag_type": "management_exit", "expected_weight": -8,
                               "label": {"event_type": "management_exit", "subject": "company", "occurred_at_company": True, "people_direction": "appointment",
                                         "role_tier": "cxo_cs_cfo_compliance"}, "reason": "deliberately wrong"}) + "\n")
    res = run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus", golden_file=str(bad))
    assert res["run"]["published"] is False and res["run"]["golden"]["failed"] == 1
    m = Corpus(tmp_path / "corpus").manifests()[0]
    assert m["published"] is False and any("golden" in w for w in m["warnings"])
    assert all(r["published"] is False for r in Corpus(tmp_path / "corpus").scores())
    good = run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus2", name="o2")
    assert good["run"]["published"] is True and good["run"]["golden"]["total"] == 37


def test_b7_5_second_run_same_day_makes_zero_llm_calls(tmp_path):
    CountingClassifier.calls = 0
    run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus", clf=CountingClassifier(), name="first")
    first = CountingClassifier.calls
    CountingClassifier.calls = 0
    res = run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus", clf=CountingClassifier(), name="second")
    assert first > 0 and CountingClassifier.calls == 0 and res["run"]["n_llm_items"] == 0
    ids = Corpus(tmp_path / "corpus").run_ids()
    assert len(ids) == 2                                            # append-only: two runs, two manifests, nothing rewritten


def test_b7_6_source_mode_before_corpus_start_is_never_pit_replay(tmp_path):
    run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus")
    c = Corpus(tmp_path / "corpus")
    assert c.source_mode_for_replay(AS_OF - timedelta(days=30)) == "archive-gdelt"
    with pytest.raises(SystemExit):                                 # nothing stored is visible that early: a replay refuses
        RP.replay(AS_OF - timedelta(days=400), None, tmp_path / "corpus", tmp_path / "r", verified=None, quiet=True)


# ============================================================== B1 / B2 / B6 / diff / golden
def test_b1_tables_are_append_only_and_hold_no_body_text(tmp_path):
    res = run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus")
    c = Corpus(tmp_path / "corpus")
    arts, evs, vers, scs = c.articles(), c.events(), c.load("corpus/verifications"), c.scores()
    assert arts and evs and vers and scs
    assert all(len(a["snippet"] or "") <= 300 for a in arts) and {a["pass"] for a in arts} <= {"sentiment", "results", "governance", "actions"}
    assert all(e["prompt_version"] == "v1" and e["model_label"] == "keyword-stub" for e in evs)
    m = c.manifests()[0]
    for k in ("source_mode", "universe_hash", "prompt_version", "rubric_version", "model_label", "model_verify", "n_articles", "n_events", "n_llm_calls", "est_cost_inr", "warnings"):
        assert k in m
    with pytest.raises(FileExistsError):
        c.write_run(m["run_id"], m, [], [], [], [])
    assert res["run"]["stamp"].startswith("supplied · as_of 2026-07-03") and res["run"]["source_mode"] == "supplied"


def test_b2_label_index_is_keyed_on_prompt_and_model(tmp_path):
    run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus")
    c = Corpus(tmp_path / "corpus")
    assert c.label_index("v1", "keyword-stub") and not c.label_index("v2", "keyword-stub") and not c.label_index("v1", "other-model")


def test_b6_workbook_and_pdf_carry_the_stamps(tmp_path):
    import openpyxl
    import pymupdf
    run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus")
    wb = openpyxl.load_workbook(tmp_path / "o" / "sentinelq_report.xlsx")
    heads = [c.value for c in wb["Scorecard"][1]]
    assert "Events" in heads and any(h.startswith("Coverage map") for h in heads)
    assert "Verification record" in [c.value for c in wb["Evidence - Governance"][1]]
    assert "Confidence" in [c.value for c in wb["Evidence - Sentiment"][1]]
    text = " ".join(pg.get_text() for pg in pymupdf.open(tmp_path / "o" / "sentinelq_scorecard.pdf"))
    assert "Run stamp" in text and "rubric v" in text and "W1:" in text
    run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus", name="o2")
    wb2 = openpyxl.load_workbook(tmp_path / "o2" / "sentinelq_report.xlsx")
    assert "Delta" in wb2.sheetnames and wb2["Delta"].max_row == 2


def test_diff_lists_drivers(tmp_path):
    run(tmp_path, bajaj_auto_year(), H, name="a")
    run(tmp_path, [i for i in bajaj_auto_year() if "ransomware" not in i.title], H, name="b")
    out = RP.diff(str(tmp_path / "a"), str(tmp_path / "b"), out=tmp_path / "DIFF.md", quiet=True)
    row = out["rows"][0]
    assert "C (only in A: investigation -10)" in row["drivers"] and (tmp_path / "DIFF.md").exists()
    assert out["governance_agreement"] == 0.0


def test_golden_set_has_37_items_and_passes():
    rep = run_golden(R)
    kinds = {}
    for x in rep["results"]:
        kinds[x["expected"]] = kinds.get(x["expected"], 0) + 1
    assert rep["total"] == 37 and kinds == {"penalise": 13, "reject": 15, "recall": 9} and rep["passed"], [x for x in rep["results"] if not x["ok"]]


def test_golden_recall_checks_the_corpus_when_it_covers_the_stock(tmp_path):
    items = [art("ANGELONE", f"Angel One update {k}: client base yoy", AS_OF - timedelta(days=k), k) for k in range(0, 365, 20)]
    run(tmp_path, items, [Holding("ANGELONE", "Angel One", "BFSI")], corpus=tmp_path / "corpus")
    rep = run_golden(R, None, Corpus(tmp_path / "corpus"))
    g36 = next(x for x in rep["results"] if x["id"] == "G36")
    assert not g36["ok"] and "not found" in g36["detail"]           # the CPO resignation is not in this corpus -> recall fails, honestly
    items.append(art("ANGELONE", "Angel One Chief Product Officer Ankit Rastogi resigns", date(2026, 5, 8), 999, "moneycontrol.com"))
    run(tmp_path, items, [Holding("ANGELONE", "Angel One", "BFSI")], corpus=tmp_path / "corpus", name="o2")
    rep = run_golden(R, None, Corpus(tmp_path / "corpus"))
    assert next(x for x in rep["results"] if x["id"] == "G36")["ok"]


def test_single_lower_tier_source_is_penalised_and_confidence_is_printed(tmp_path):
    """Google News collapses syndication, so one link per event is normal: a gate-verified event keeps its penalty; confidence is printed."""
    items = [art("X", f"X Corp update {k}: volumes yoy", AS_OF - timedelta(days=k), k) for k in range(0, 365, 15)]
    items.append(art("X", "X Corp faces regulator probe over disclosures", AS_OF - timedelta(days=5), 700, "unknown-blog.in"))
    s = run(tmp_path, items, [Holding("X", "X Corp", "Autos")])["scores"][0]
    assert s.governance_score == 80 and "verify against the filing" in s.governance_penalties[0]["basis"]
    assert s.governance_penalties[0]["confidence"] < cfg.CONF_UNPENALISED_BELOW


def test_orderly_succession_with_named_successor_is_not_an_exit():
    from sentinelq.models import Label, LabelledItem
    from sentinelq.score import gate
    mk = lambda t: LabelledItem(RawItem("N", "news", t, "", "https://reuters.com/x", "2025-12-10", "reuters.com"),
                                Label("management_exit", -1, "CFO change", True, subject="company", occurred_at_company=True,
                                      people_direction="unplanned_exit", role_tier="cxo_cs_cfo_compliance"))
    assert gate(mk("Nestle India CFO Svetlana Boldina to step down, successor to be named"), R)[0] is None
    assert gate(mk("Cummins India MD to step down; Shveta Arya to take over"), R)[0] is None
    assert gate(mk("Nestle India CFO resigns with immediate effect"), R)[0] == -8
    assert gate(mk("Zen Technologies CFO resigns amid audit queries"), R)[0] == -8


def test_signal_event_aggregation_reaches_plus_two_and_offset_caps_it():
    from sentinelq.models import Label, LabelledItem
    from sentinelq.score import sentiment_terms
    mk = lambda i, et, sent: LabelledItem(RawItem("B", "news", f"h{i}", "", f"https://reuters.com/{i}", (AS_OF - timedelta(days=i)).isoformat(), "reuters.com",
                                                   confidence=0.5, pass_="results" if et.startswith("earn") else "sentiment"), Label(et, sent, "", False))
    items = [mk(i, "earnings_beat", 2) for i in range(4)] + [mk(10 + i, "analyst_action", 0) for i in range(30)]
    terms, mean, rule = sentiment_terms(items, AS_OF, R)
    assert mean == 2.0 and rule.startswith("signal-event mean over 4")                 # thirty neutral notes no longer dilute four beats
    assert sum(1 for t in terms if not t[2]) == 30
    items.append(mk(50, "earnings_miss", -1))
    terms, mean, rule = sentiment_terms(items, AS_OF, R)
    assert mean == 1.49 and "+2 withheld" in rule                                      # the rubric's own anchor: +2 needs no material offset
    few = [mk(0, "earnings_beat", 2)] + [mk(10 + i, "analyst_action", 0) for i in range(5)]
    assert sentiment_terms(few, AS_OF, R)[2].startswith("all-event")                   # below min_signal_events: fall back


def test_standing_verified_convention_applies_at_any_as_of(tmp_path):
    from sentinelq.verified import load_verified
    csv = tmp_path / "v.csv"
    csv.write_text("symbol,date,headline,event_type,sentiment,subject,severity,action_stage,people_direction,role_tier,amount_inr_cr,event_key,penalty_override,status,valid_as_of,source_ref,url,note,standing\n"
                   "BOSCHLTD,,MNC structural discount,mnc_structural_discount,0,company,procedural,n/a,n/a,n/a,,bosch-mnc,-5,convention,2026-07-03,recon,,,\n"
                   "NAM-INDIA,,CBI probe open,investigation,-1,company,material,investigation,n/a,n/a,,nippon-cbi,-10,verified,2026-07-03,BW,,,\n")
    vit, skipped = load_verified(csv, date(2026, 10, 5), 365, {"BOSCHLTD", "NAM-INDIA"})
    assert [v.item.symbol for v in vit] == ["BOSCHLTD"] and skipped[0]["symbol"] == "NAM-INDIA"
    real, _ = load_verified(ROOT / "portfolio" / "verified_events.csv", date(2026, 10, 5), 365, {"BOSCHLTD"})
    assert any("mnc" in (v.label.event_key or "") for v in real)


def test_corpus_union_restores_a_governance_event_missed_by_a_later_run(tmp_path):
    base = bajaj_auto_year()
    run(tmp_path, base, H, corpus=tmp_path / "corpus", name="first")                       # the ransomware probe is on record
    without = [i for i in base if "ransomware" not in i.title]
    res = run(tmp_path, without, H, corpus=tmp_path / "corpus", name="second")                # this run's feed lost it
    s = res["scores"][0]
    assert res["run"]["corpus_union"]["events"] >= 1 and s.governance_score == 90
    inv = next(p for p in s.governance_penalties if p["event_type"] == "investigation")
    assert any(li.item.origin.startswith("corpus:") for li in res["kept"]["BAJAJ-AUTO"] if li.label.event_type == "investigation")
    res2 = run(tmp_path, without, H, corpus=tmp_path / "corpus", name="third", corpus_union=False)
    assert res2["scores"][0].governance_score == 100


def test_results_quota_keeps_one_print_plus_other_results_type_events_per_quarter(tmp_path):
    base = [art("B", f"Bajaj Auto update {k}: volumes yoy", AS_OF - timedelta(days=k), k) for k in range(0, 365, 9)]
    q = AS_OF - timedelta(days=20)
    extra = [art("B", f"Bajaj Auto Q1 results: profit rises {i}", q - timedelta(days=8 * i), 600 + i) for i in range(3)]          # three prints, same quarter, > 2 days apart
    extra += [art("B", f"Bajaj Auto guidance raised on exports {i}", q - timedelta(days=30 + 6 * i), 700 + i) for i in range(5)]   # five other results-type items
    p = Pipeline(R, Fixed(base), None, None, KeywordClassifier(), tmp_path / "o", None, AS_OF, governance_pass=False, fundamentals_pass=True,
                 stratified=True, corpus_dir=None, disk_cache=False)
    p.news.fetch_fundamentals = lambda h, s, e: extra                                   # the results pass is a separate search
    res = p.run([Holding("B", "Bajaj Auto", "Autos")])
    reasons = [d.reason for d in res["dropped"] if d.reason.startswith("results_quota")]
    assert reasons.count("results_quota_one_print_per_quarter") >= 1 and reasons.count("results_quota_other_per_quarter") >= 1
    kept_res = [li for li in res["kept"]["B"] if li.item.pass_ == "results"]
    assert 1 <= sum(1 for li in kept_res if "results" in li.item.title) <= 1 and sum(1 for li in kept_res if "guidance" in li.item.title) <= cfg.RESULTS_OTHER_PER_QUARTER


def test_cli_dispatches_the_new_commands(tmp_path, capsys):
    from sentinelq import cli
    with pytest.raises(SystemExit) as e:
        cli.main(["golden"])
    assert e.value.code == 0 and "37/37" in capsys.readouterr().out


def test_governance_cap_is_stratified_so_recent_routine_headlines_cannot_crowd_out_an_old_order():
    from sentinelq.ingest.select import select_governance
    recent = [art("X", f"X Corp appoints VP {i} as head", AS_OF - timedelta(days=i), i) for i in range(1, 41)]          # 40 routine people items in W1/W2
    old = art("X", "SEBI order against X Corp over disclosure lapses", AS_OF - timedelta(days=300), 900, "reuters.com")   # one real order in W4
    legacy = select_governance(recent + [old], [], 30, ["x corp"])
    assert old not in legacy                                                   # newest-first: the order never makes the 30
    strat = select_governance(recent + [old], [], 30, ["x corp"], as_of=AS_OF)
    assert old in strat and len(strat) == 30 and old.window == 4


def test_learn_proposes_clean_inflation(tmp_path):
    from sentinelq import learn as L
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    sc = [{"symbol": f"S{i}", "governance_score": 100, "governance_label": "Clean", "governance_penalties": [], "company_sentiment": 1,
           "sentiment_note": "", "relevant_articles": 30, "evidence_span_days": 300, "governance_ignored": [{"why": "x"}]} for i in range(12)]
    (run_dir / "scores.json").write_text(json.dumps(sc))
    (run_dir / "evidence.json").write_text("[]")
    (run_dir / "run_meta.json").write_text(json.dumps({"as_of": "2026-10-05"}))
    kinds = {l["kind"] for l in L.detect(run_dir)}
    assert "clean_inflation" in kinds


def test_label_split_observation_is_deterministic_and_compares_to_previous_run(tmp_path):
    run(tmp_path, bajaj_auto_year(), H, corpus=tmp_path / "corpus", name="a")
    res = run(tmp_path, [i for i in bajaj_auto_year() if "ransomware" not in i.title], H, corpus=tmp_path / "corpus", name="b", corpus_union=False)
    obs = [o for o in res["observations"] if o["title"].startswith("Label split")]
    assert len(obs) == 1 and "Previous run (" in obs[0]["body"] and "BAJAJ-AUTO (governance 90" in obs[0]["body"]
    assert "Clean 1 / Watch 0 / Flag 0" in obs[0]["body"] and "Watch 1" in obs[0]["body"].split("Previous run")[1]
